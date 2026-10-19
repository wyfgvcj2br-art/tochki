import datetime,io,re,sys,tempfile,unittest,zipfile
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'app'))
import openpyxl
from engine import analyze,select,generate,clean_name,selection_inputs
def xlsx(rows):
    w=openpyxl.Workbook()
    for row in rows:w.active.append(row)
    b=io.BytesIO();w.save(b);return b.getvalue()
def with_dimension(blob,dimension):
    output=io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(blob)) as source,zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as target:
        for name in source.namelist():
            value=source.read(name)
            if name=='xl/worksheets/sheet1.xml':value=re.sub(rb'<dimension\b[^>]*/>',f'<dimension ref="{dimension}"/>'.encode(),value)
            target.writestr(name,value)
    return output.getvalue()
HEADER=['Наименование ТТ','Дата','Заказы в статусе Успешно завершен','Общее кол-во заказов, просроченных по сборке','Помощник']
TEST_TMP=Path(__file__).resolve().parents[1]/'.cache/tests'
TEST_TMP.mkdir(parents=True,exist_ok=True)
class EngineTests(unittest.TestCase):
    def setUp(self):
        self.daily=xlsx([HEADER,['9001_Демо','2026-01-05',100,5,'Помощник Демо'],['9001_Демо','2026-01-06',300,30,'Помощник Демо']])
        self.appeals=xlsx([['Сборщик','ТТ','абс','Дни'],['Тестовый Сотрудник 123(456)','9001_Демо',2,'2026-01-05'],['()','9001_Демо',1,'2026-01-06']])
        self.options={'codes':['9001'],'excludeUnknown':True,'journal':True,'outline':True,'good':6,'bad':10,'mode':'combined'}
    def data(self):return analyze([('daily.xlsx',self.daily),('appeals.xlsx',self.appeals)])
    def test_aggregation_and_exports(self):
        data=self.data();self.assertFalse(data['errors']);groups,dates,apps=select(data,self.options)
        self.assertEqual((groups[0]['orders'],groups[0]['late'],len(apps)),(400,35,1))
        self.assertEqual(apps[0]['name'],'Тестовый Сотрудник')
        with tempfile.TemporaryDirectory(dir=TEST_TMP) as tmp:
            audit=generate(data,self.options,tmp);self.assertEqual(audit['weeklyPct'],35/400);self.assertEqual(audit['appeals'],2)
            w=openpyxl.load_workbook(Path(tmp)/'Отчёт.xlsx');self.assertEqual(w.sheetnames,['По торговым точкам','Журнал обращений','Структура']);self.assertFalse(any(r.hidden for r in w.worksheets[0].row_dimensions.values()))
            self.assertEqual(len(list((Path(tmp)/'Изображения').glob('*.png'))),1)
            self.assertIsNone(zipfile.ZipFile(Path(tmp)/'Все результаты.zip').testzip())
    def test_file_deduplication(self):
        data=analyze([('a.xlsx',self.daily),('copy.xlsx',self.daily)]);self.assertEqual(len(data['daily']),2);self.assertTrue(data['warnings'])
    def test_conflicting_rows_block_generation(self):
        changed=xlsx([HEADER,['9001_Демо','2026-01-05',200,10,'Демо']]);data=analyze([('a.xlsx',self.daily),('b.xlsx',changed)]);self.assertTrue(data['errors'])
        with tempfile.TemporaryDirectory(dir=TEST_TMP) as tmp:
            with self.assertRaises(ValueError):generate(data,self.options,tmp)
    def test_unmatched_appeal_is_visible(self):
        data=self.data();data['daily'][1]['date']='2026-01-07';data['appeals'][0]['date']='2026-01-06'
        with self.assertRaises(ValueError):select(data,self.options)
    def test_appeals_outside_daily_period_are_reported_and_excluded(self):
        data=self.data()
        for date,count in [('2025-07-17',3),('2026-01-07',4)]:
            data['appeals'].append({**data['appeals'][0],'date':date,'count':count})
        data['appeals'].append({**data['appeals'][1],'date':'2025-07-17','count':5})
        groups,dates,apps=select(data,self.options)
        self.assertEqual(dates,['2026-01-05','2026-01-06']);self.assertEqual(sum(a['count'] for a in apps),2)
        scope=selection_inputs(data,self.options)[2]
        self.assertEqual(scope['outsidePeriodRecords'],3);self.assertEqual(scope['outsidePeriodAppeals'],12)
        self.assertEqual(scope['excludedUnknown'],1);self.assertEqual(scope['excludedUnknownAppeals'],1)
        self.assertIn('05.01.2026 — 06.01.2026',scope['warnings'][-1])
        with tempfile.TemporaryDirectory(dir=TEST_TMP) as tmp:
            audit=generate(data,self.options,tmp)
            self.assertEqual(audit['outsidePeriodRecords'],3);self.assertEqual(audit['appeals'],2)
            self.assertEqual(audit['period'],['2026-01-05','2026-01-06']);self.assertEqual(audit['excludedUnknown'],1)
    def test_period_boundaries_include_appeals(self):
        data=self.data();data['appeals'][1]['unknown']=False;data['appeals'][1]['name']='Второй Сотрудник'
        groups,dates,apps=select(data,self.options)
        self.assertEqual(len(apps),2);self.assertEqual(sum(a['count'] for a in apps),3)
    def test_appeals_only_does_not_limit_to_daily_period(self):
        data=self.data();data['appeals'][0]['date']='2025-07-17'
        groups,dates,apps=select(data,{**self.options,'mode':'appeals'})
        self.assertEqual(dates,['2025-07-17']);self.assertEqual(len(apps),1);self.assertEqual(groups[0]['days'],[])
        self.assertEqual(selection_inputs(data,{**self.options,'mode':'appeals'})[2]['outsidePeriodRecords'],0)
    def test_missing_date_for_store_inside_period_still_blocks(self):
        data=self.data();data['daily'].append({**data['daily'][0],'code':'9002','tt':'9002_Демо','date':'2026-01-07'})
        data['appeals'][0]['date']='2026-01-07'
        with self.assertRaises(ValueError):select(data,{**self.options,'codes':['9001','9002']})
    def test_unknown_is_opt_in(self):
        data=self.data();self.assertEqual(len(select(data,self.options)[2]),1);self.assertEqual(len(select(data,{**self.options,'excludeUnknown':False})[2]),2)
    def test_late_only_ignores_unmatched_appeals(self):
        data=self.data();data['daily']=data['daily'][1:];self.assertEqual(select(data,{**self.options,'mode':'late'})[2],[])
    def test_zero_and_appeals_only(self):
        data=analyze([('zero.xlsx',xlsx([HEADER,['9001_Демо','2026-01-05',0,0,'Демо']]))])
        with tempfile.TemporaryDirectory(dir=TEST_TMP) as tmp:self.assertIsNone(generate(data,self.options,tmp)['weeklyPct'])
        data=analyze([('appeals.xlsx',self.appeals)])
        with tempfile.TemporaryDirectory(dir=TEST_TMP) as tmp:self.assertEqual(generate(data,{**self.options,'mode':'appeals'},tmp)['appeals'],2)
    def test_clean_names(self):self.assertEqual(clean_name('Тестовый Сотрудник 123(456)'),'Тестовый Сотрудник')
    def test_bi_appeals_with_incorrect_dimensions(self):
        blob=xlsx([['Примененные фильтры: отчёт за неделю'],[],['Сборщик','ТТ','абс','Дни'],['Тестовый Сотрудник(123)','9001_Демо',2,'2026-01-05'],['()','9001_Демо',1,'2026-01-06']])
        for dimension in ['A1','A1:D4','A1:A5']:
            with self.subTest(dimension=dimension):
                data=analyze([('bi-export.xlsx',with_dimension(blob,dimension))])
                self.assertFalse(data['errors']);self.assertEqual(len(data['appeals']),2)
                self.assertEqual([(r['date'],r['count']) for r in data['appeals']],[('2026-01-05',2),('2026-01-06',1)])
                self.assertEqual(data['unknownCount'],1);self.assertEqual(data['sources'][0]['header'],3)
    def test_daily_with_incorrect_dimensions(self):
        data=analyze([('daily.xlsx',with_dimension(self.daily,'A1'))])
        self.assertFalse(data['errors']);self.assertEqual(len(data['daily']),2)
        self.assertEqual(sum(r['orders'] for r in data['daily']),400);self.assertEqual(sum(r['late'] for r in data['daily']),35)
if __name__=='__main__':unittest.main()
