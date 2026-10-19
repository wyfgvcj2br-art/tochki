import datetime as dt
import io
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

import openpyxl
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'app'))
from engine import analyze,generate
from details import detail_selection


def book(rows):
    w=openpyxl.Workbook()
    for row in rows:w.active.append(row)
    output=io.BytesIO();w.save(output);return output.getvalue()


DAILY=['Наименование ТТ','Дата','Заказы в статусе Успешно завершен','Общее кол-во заказов, просроченных по сборке']
PEOPLE=['Наименование ТТ','date_collected','Номер сборщика','Сборщик','Свойство сборщика','Заказы в статусе Успешно завершен','Общее кол-во заказов, просроченных по сборке','Общая доля заказов, просроченных по сборке']
TMP=Path(__file__).resolve().parents[1]/'.cache/tests';TMP.mkdir(parents=True,exist_ok=True)


class DetailTests(unittest.TestCase):
    def setUp(self):
        self.daily=[DAILY,['9001_Демо','2026-01-05',100,5],['9001_Демо','2026-01-06',100,7]]
        self.people=[PEOPLE,['9001_Демо',dt.datetime(2026,1,6),'1','Первый Сборщик(99)','Продавец',40,7,7/40],['9001_Демо',dt.datetime(2026,1,6),'2','N/A','N/A',60,0,0]]
        self.opts={'codes':['9001'],'mode':'details','excludeUnknown':True}
    def data(self,daily=None,people=None):
        return analyze([('daily.xlsx',book(daily or self.daily)),('people.xlsx',book(people or self.people))])
    def test_separate_grain_and_inclusive_store_threshold(self):
        data=self.data();self.assertFalse(data['errors']);self.assertEqual(len(data['daily']),2);self.assertEqual(len(data['collectors']),2)
        p=detail_selection(data,self.opts);self.assertEqual((p['badDays'],p['collectorCount'],p['orders'],p['late']),(1,2,100,7))
        self.assertEqual(p['groups'][0]['days'][0]['date'],'2026-01-06')
        self.assertEqual(p['groups'][0]['days'][0]['collectors'][0]['name'],'Первый Сборщик')
        self.assertEqual(p['unknownCollectors'],1);self.assertTrue(p['warnings'])
    def test_only_store_rate_selects_days_not_individual_rate(self):
        data=self.data();data['daily'][1]['late']=5;data['daily'][1]['pct']=.05
        p=detail_selection(data,self.opts);self.assertEqual(p['groups'],[]);self.assertFalse(p['canGenerate'])
        with tempfile.TemporaryDirectory(dir=TMP) as tmp:
            with self.assertRaisesRegex(ValueError,'не требуется'):generate(data,self.opts,tmp)
    def test_exactly_six_percent_is_included(self):
        data=self.data();data['daily'][1]['late']=6;data['daily'][1]['pct']=.06;data['collectors'][0]['late']=6
        result=detail_selection(data,self.opts);self.assertEqual(result['badDays'],1)
        self.assertEqual(result['groups'][0]['days'][0]['pct'],.06)
    def test_before_rounding_threshold(self):
        daily=[DAILY,['9001_Демо','2026-01-05',100000,5999]]
        people=[PEOPLE,['9001_Демо','2026-01-05','1','Сборщик','Продавец',100000,5999,.05999]]
        self.assertEqual(detail_selection(self.data(daily,people),self.opts)['badDays'],0)
    def test_missing_day_blocks_but_differences_warn_without_changing_collectors(self):
        data=self.data();data['collectors'][0]['date']='2026-01-07';data['collectors'][1]['date']='2026-01-07'
        with self.assertRaisesRegex(ValueError,'Нет детализации'):detail_selection(data,self.opts)
        data=self.data();data['collectors'][0]['orders']=39
        preview=detail_selection(data,self.opts)
        self.assertEqual(preview['mismatchCount'],1);self.assertTrue(preview['warnings']);self.assertTrue(preview['canGenerate'])
        self.assertEqual((preview['orders'],preview['late'],preview['storeOrders']),(99,7,100))
        self.assertEqual(preview['groups'][0]['days'][0]['collectors'][0]['orders'],39)
        self.assertEqual(preview['groups'][0]['days'][0]['pct'],7/99)
    def test_store_selects_day_but_collector_summary_below_threshold_is_preserved(self):
        data=self.data();data['collectors'][0]['late']=5
        p=detail_selection(data,self.opts);self.assertEqual(p['badDays'],1)
        self.assertEqual(p['groups'][0]['days'][0]['pct'],.05);self.assertEqual(p['groups'][0]['days'][0]['storePct'],.07)
        self.assertEqual(p['mismatchCount'],1)
    def test_adjacent_day_shift_only_warns_and_keeps_both_source_days(self):
        daily=[DAILY,['9001_Демо','2026-01-05',37,5],['9001_Демо','2026-01-06',42,3]]
        people=[PEOPLE,['9001_Демо','2026-01-05','1','Первый','Продавец',36,4,4/36],['9001_Демо','2026-01-06','1','Первый','Продавец',43,4,4/43]]
        data=self.data(daily,people);p=detail_selection(data,self.opts)
        self.assertEqual(p['mismatchCount'],2)
        self.assertEqual([(d['orders'],d['late'],d['pct']) for d in p['groups'][0]['days']],[(36,4,4/36),(43,4,4/43)])
        self.assertEqual((p['orders'],p['late']),(79,8))
        with tempfile.TemporaryDirectory(dir=TMP) as tmp:
            audit=generate(data,self.opts,tmp);self.assertEqual(audit['mismatchCount'],2)
            wb=openpyxl.load_workbook(Path(tmp)/'Отчёт.xlsx');values=[r for r in wb.active.values if isinstance(r[0],dt.datetime)]
            self.assertEqual([(r[1],r[2],r[3]) for r in values],[(36,4,4/36),(43,4,4/43)])
    def test_maximum_seven_calendar_days_without_silent_truncation(self):
        daily=[DAILY,['9001_Демо','2026-01-05',100,7],['9001_Демо','2026-01-12',100,7]]
        with self.assertRaisesRegex(ValueError,'7 календарных'):detail_selection(self.data(daily),self.opts)
    def test_one_image_for_seven_days_and_second_store(self):
        daily=[DAILY];people=[PEOPLE]
        for i in range(7):
            day=f'2026-01-{5+i:02}'
            daily.append(['9001_Демо',day,100,7]);people.append(['9001_Демо',day,'1','Сборщик','Продавец',100,7,.07])
        daily.append(['9002_Другая','2026-01-05',100,8]);people.append(['9002_Другая','2026-01-05','2','Второй','Продавец',100,8,.08])
        with tempfile.TemporaryDirectory(dir=TMP) as tmp:
            result=generate(self.data(daily,people),{**self.opts,'codes':['9001','9002']},tmp)
            self.assertEqual((result['stores'],result['badDays']),(2,8));self.assertEqual(len(list((Path(tmp)/'Изображения').glob('*.png'))),2)
            wb=openpyxl.load_workbook(Path(tmp)/'Отчёт.xlsx');self.assertEqual(wb.sheetnames,['Детализация сборщиков'])
            self.assertIsNone(zipfile.ZipFile(Path(tmp)/'Все результаты.zip').testzip())
    def test_store_alias_and_name_variant_are_merged_by_person_id(self):
        daily=[DAILY,['9001М_Демо','2026-01-06',100,7]]
        people=[PEOPLE,['9001М_Демо','2026-01-06','1','Сборщик','Продавец',90,7,7/90],['9001Ш_Демо','2026-01-06','1','N/A','N/A',10,0,0]]
        p=detail_selection(self.data(daily,people),self.opts);person=p['groups'][0]['days'][0]['collectors'][0]
        self.assertEqual(p['collectorCount'],1);self.assertEqual((person['name'],person['orders'],person['late'],person['pct']),('Сборщик',100,7,.07))
    def test_duplicate_file_and_duplicate_rows_are_not_double_counted(self):
        blob=book(self.people);data=analyze([('d.xlsx',book(self.daily)),('p.xlsx',blob),('copy.xlsx',blob)])
        self.assertEqual(detail_selection(data,self.opts)['orders'],100)
        data=self.data(people=self.people+[self.people[1]])
        self.assertFalse(data['errors']);self.assertEqual(len(data['collectors']),2);self.assertTrue(data['warnings'])
    def test_conflicting_person_rows_and_percentages_rejected(self):
        changed=self.people[1].copy();changed[-3:]=[41,7,7/41]
        self.assertTrue(self.data(people=self.people+[changed])['errors'])
        changed=self.people[1].copy();changed[-1]=.9
        self.assertTrue(self.data(people=[PEOPLE,changed])['errors'])
    def test_missing_required_sources_and_zero_orders(self):
        data=analyze([('p.xlsx',book(self.people))]);self.assertFalse(data['errors'])
        with self.assertRaisesRegex(ValueError,'по магазинам'):detail_selection(data,self.opts)
        data=analyze([('d.xlsx',book(self.daily))])
        with self.assertRaisesRegex(ValueError,'по сборщикам'):detail_selection(data,self.opts)
        extra=['9001_Демо','2026-01-06','3','N/A','N/A',0,0,'N/A']
        data=self.data(people=self.people+[extra]);self.assertFalse(data['errors']);self.assertEqual(detail_selection(data,self.opts)['collectorCount'],3)


if __name__=='__main__':unittest.main()
