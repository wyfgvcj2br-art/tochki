"""Collector details for store days whose pooled overdue rate is >=6%."""
from __future__ import annotations
import collections
import datetime as dt
import json
import math
import re
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageDraw

from engine import DEFAULT_CODES, BLUE, font, n, pct
from collector_names import collector_id

LIMIT=.06


def detail_color(value):
    return '#EDF2F6' if value is None else '#C6EFCE' if value<LIMIT else '#F08080' if value>.10 else '#FCE0E0'


def detail_selection(data,opts):
    if data['errors']:raise ValueError('Сначала устраните ошибки в загруженных данных.')
    selected=set(opts.get('codes',DEFAULT_CODES))
    days=sorted((r for r in data['daily'] if r['code'] in selected),key=lambda r:(r['code'],r['date']))
    if not days:raise ValueError('Для детализации добавьте «Итоговая таблица по магазинам по дням» по выбранным ТТ: по ней определяются дни с просроком от 6% включительно.')
    dates=sorted({r['date'] for r in days})
    if (dt.date.fromisoformat(dates[-1])-dt.date.fromisoformat(dates[0])).days>=7:
        raise ValueError('Детализация: загрузите выгрузку по магазинам максимум за 7 календарных дней, чтобы на каждую ТТ получился один PNG.')
    if not data.get('collectors'):raise ValueError('Добавьте отдельный XLSX «Итоговая таблица по сборщикам по дням».')
    bad=[r for r in days if r['orders'] and r['late']/r['orders']>=LIMIT]
    indexed=collections.defaultdict(list)
    for row in data['collectors']:indexed[row['code'],row['date']].append(row)
    grouped=collections.defaultdict(list);unknown_rows=0;source_rows=0;mismatches=[];hidden=[]
    for daily in bad:
        key=daily['code'],daily['date'];raw=indexed[key]
        label=f'{daily["tt"]}, {dt.date.fromisoformat(daily["date"]):%d.%m.%Y}'
        if not raw:raise ValueError(f'Нет детализации сборщиков для {label}. Добавьте выгрузку за этот день.')
        hidden.extend(dict(r) for r in raw if collector_id(r.get('collectorId'))=='0')
        source_rows+=len(raw)
        raw=[r for r in raw if collector_id(r.get('collectorId'))!='0']
        orders=sum(r['orders'] for r in raw);late=sum(r['late'] for r in raw)
        differs=not math.isclose(orders,daily['orders'],rel_tol=1e-10,abs_tol=1e-9) or not math.isclose(late,daily['late'],rel_tol=1e-10,abs_tol=1e-9)
        different_date_fields=daily.get('dateField')=='Дата и время начала слота' and all(r.get('dateField')=='date_collected' for r in raw)
        if differs:
            mismatches.append({'code':daily['code'],'tt':daily['tt'],'date':daily['date'],'storeOrders':daily['orders'],'storeLate':daily['late'],'collectorOrders':orders,'collectorLate':late,'orderDelta':daily['orders']-orders,'lateDelta':daily['late']-late,'storeDateField':daily.get('dateField','Дата'),'collectorDateFields':sorted({r.get('dateField','Дата') for r in raw})})
        people={}
        for row in raw:
            # Store aliases and source role/name variants may split one person's day.
            identity=('id',row['collectorId']) if row['collectorId'] else ('name',row['name'])
            person=people.setdefault(identity,{'name':'','nameFromRegistry':False,'collectorId':row['collectorId'],'orders':0,'late':0,'sources':[]})
            if row['name']:
                mapped=bool(row.get('nameFromRegistry'))
                if person['name'] and person['name']!=row['name'] and not mapped and not person['nameFromRegistry']:raise ValueError(f'У одного номера сборщика разные ФИО: {label}. Проверьте исходную выгрузку.')
                if not person['name'] or not mapped:person.update(name=row['name'],nameFromRegistry=mapped)
            person['orders']+=row['orders'];person['late']+=row['late'];person['sources'].append(row['source'])
        rows=list(people.values())
        for person in rows:
            person['pct']=person['late']/person['orders'] if person['orders'] else None
            if not person['name']:unknown_rows+=1
        rows.sort(key=lambda r:(-r['late'],-(r['pct'] or 0),r['name'].casefold(),r['collectorId']))
        # Store counts select the day only. Every displayed measure comes from collectors.
        grouped[daily['code']].append({**daily,'storeOrders':daily['orders'],'storeLate':daily['late'],'storePct':daily['late']/daily['orders'],'orders':orders,'late':late,'pct':late/orders if orders else None,'collectors':rows,'sourceRows':len(raw),'totalsDiffer':differs,'differentDateFields':different_date_fields})
    groups=[]
    for code,records in sorted(grouped.items()):
        differs=any(r['totalsDiffer'] for r in records)
        note='Показатели — из выгрузки сборщиков; дни выбраны по просроку ТТ.'
        if any(collector_id(r.get('collectorId'))=='0' and r['code']==code for r in hidden):note='ID 0 исключён из строк и итогов; дни — по исходному просроку ТТ.'
        elif differs:note='Есть расхождения с итогами ТТ. Показатели — из выгрузки сборщиков.'
        groups.append({'code':code,'tt':records[0]['tt'],'days':records,'orders':sum(r['orders'] for r in records),'late':sum(r['late'] for r in records),'storeOrders':sum(r['storeOrders'] for r in records),'storeLate':sum(r['storeLate'] for r in records),'totalsDiffer':differs,'dateNote':note,'appeals':[]})
    warnings=list(data['warnings'])
    if hidden:warnings.append(f'ID 0 скрыт из детализации и её итогов: строк — {len(hidden)}, заказов — {sum(r["orders"] for r in hidden):g}, просроченных — {sum(r["late"] for r in hidden):g}. Порог дней по исходной таблице ТТ не изменён.')
    for mismatch in mismatches:
        label=f'{mismatch["tt"]}, {dt.date.fromisoformat(mismatch["date"]):%d.%m.%Y}'
        warnings.append(f'{label}: ТТ — {mismatch["storeOrders"]:g} заказов / {mismatch["storeLate"]:g} просроченных; сборщики — {mismatch["collectorOrders"]:g} / {mismatch["collectorLate"]:g}. Детализация взята из таблицы сборщиков без изменения чисел. Причина расхождения по агрегированной выгрузке не подтверждена.')
    normal=sorted({r['code'] for r in days}-set(grouped))
    if normal:warnings.append('Нет дней с просроком от 6% включительно; PNG не требуется для ТТ: '+', '.join(normal)+'.')
    if unknown_rows:warnings.append(f'Сборщиков без ФИО в детализации: {unknown_rows}. Их заказы сохранены; ФИО можно привязать к ID в справочнике.')
    if not groups:warnings.append('За выбранный период ни у одной выбранной ТТ нет дней с общим просроком от 6% включительно.')
    return {'mode':'details','nameBindingsUsed':data.get('nameBindingsUsed',[]),'nameRegistryRevision':data.get('nameRegistryRevision'),'hiddenIdZero':{'rows':len(hidden),'orders':sum(r['orders'] for r in hidden),'late':sum(r['late'] for r in hidden)},'measureSource':'collectors','groups':groups,'dates':dates,'period':[dates[0],dates[-1]],'threshold':LIMIT,'thresholdInclusive':True,'badDays':len(bad),'collectorCount':sum(len(d['collectors']) for g in groups for d in g['days']),'collectorSourceRows':source_rows,'unknownCollectors':unknown_rows,'normalCodes':normal,'mismatchCount':len(mismatches),'mismatches':mismatches,'appeals':0,'orders':sum(g['orders'] for g in groups),'late':sum(g['late'] for g in groups),'storeOrders':sum(g['storeOrders'] for g in groups),'storeLate':sum(g['storeLate'] for g in groups),'warnings':warnings,'canGenerate':bool(groups)}


def detail_rows(group):
    rows=[]
    for day in group['days']:
        rows.append(([dt.date.fromisoformat(day['date']).strftime('%d.%m.%Y'),n(day['orders']),n(day['late']),pct(day['pct'])],day['pct'],'day'))
        for p in day['collectors']:rows.append(([p['name'] or 'ФИО не указано',n(p['orders']),n(p['late']),pct(p['pct'])],p['pct'],'person'))
    ratio=group['late']/group['orders'] if group['orders'] else None
    rows.append((['Итого по сборщикам за эти дни',n(group['orders']),n(group['late']),pct(ratio)],ratio,'total'))
    return rows


def detail_picture(group,dates,path):
    rows=detail_rows(group);width=1500;edges=(0,820,1040,1260,1500);header=193;rh=50
    im=Image.new('RGB',(width,header+rh*len(rows)),'white');draw=ImageDraw.Draw(im)
    draw.text((20,16),group['tt'],font=font(30,True),fill='#243E52')
    draw.text((20,62),f'{dt.date.fromisoformat(dates[0]):%d.%m.%Y} — {dt.date.fromisoformat(dates[-1]):%d.%m.%Y}',font=font(20),fill='#607484')
    draw.text((20,96),f'Детализация сборщиков · просрок ТТ ≥ 6% · дней: {len(group["days"])}',font=font(20),fill='#607484')
    draw.rectangle((0,133,width,header),fill='#'+BLUE)
    for i,label in enumerate(('Дата / Сборщик','Заказы УЗ, шт.','Просрочено, шт.','Общий просрок')):
        draw.text(((edges[i]+edges[i+1])/2,163),label,font=font(20,True),fill='white',anchor='mm')
    for j,(values,ratio,kind) in enumerate(rows):
        y=header+j*rh;bold=kind!='person';fill='#DCE6ED' if bold else ('white' if j%2 else '#F7FAFC')
        draw.rectangle((0,y,width,y+rh),fill=fill)
        draw.rectangle((1260,y,width,y+rh),fill=detail_color(ratio))
        for col,value in enumerate(values):
            face=font(21,bold or col==3)
            if col==0:
                while draw.textlength(value,font=face)>775 and face.size>14:face=font(face.size-1,bold)
                draw.text((20,y+25),value,font=face,fill='#243D50',anchor='lm')
            else:draw.text(((edges[col]+edges[col+1])/2,y+25),value,font=face,fill='#243D50',anchor='mm')
        draw.line((0,y+rh-1,width,y+rh-1),fill='#DCE4EA')
    im.save(path)


def generate_details(data,opts,folder):
    preview=detail_selection(data,opts)
    if not preview['canGenerate']:raise ValueError('Дней с общим просроком ТТ от 6% включительно нет. Детализация не требуется.')
    folder=Path(folder);folder.mkdir(parents=True,exist_ok=True);images=folder/'Изображения';images.mkdir(exist_ok=True)
    wb=openpyxl.Workbook();ws=wb.active;ws.title='Детализация сборщиков';ws.sheet_view.showGridLines=False
    ws.column_dimensions['A'].width=70
    for col in 'BCD':ws.column_dimensions[col].width=23
    for group in preview['groups']:
        ws.append([group['tt']]);ws.cell(ws.max_row,1).font=Font(name='Arial',size=15,bold=True,color=BLUE)
        ws.append([f'{dt.date.fromisoformat(preview["dates"][0]):%d.%m.%Y} — {dt.date.fromisoformat(preview["dates"][-1]):%d.%m.%Y} · Только дни с просроком ТТ ≥ 6%'])
        ws.append(['Дата / Сборщик','Заказы УЗ, шт.','Просрочено, шт.','Общий просрок'])
        for c in ws[ws.max_row]:c.fill=PatternFill('solid',fgColor=BLUE);c.font=Font(name='Arial',bold=True,color='FFFFFF')
        refs=[]
        for day in group['days']:
            start=ws.max_row+1;refs.append(start);ws.append([dt.date.fromisoformat(day['date']),day['orders'],day['late'],day['pct']]);ws.cell(start,1).number_format='dd.mm.yyyy'
            for p in day['collectors']:ws.append([p['name'] or 'ФИО не указано',p['orders'],p['late'],p['pct']])
            for c in ws[start]:c.font=Font(name='Arial',bold=True,size=11);c.fill=PatternFill('solid',fgColor='DCE6ED')
        total=ws.max_row+1
        ws.append(['Итого по сборщикам за эти дни','=SUM('+','.join(f'B{r}' for r in refs)+')','=SUM('+','.join(f'C{r}' for r in refs)+')',f'=IF(B{total}=0,"",C{total}/B{total})'])
        for c in ws[total]:c.font=Font(name='Arial',bold=True,size=11);c.fill=PatternFill('solid',fgColor='DCE6ED')
        ws.append([]);ws.append([])
        detail_picture(group,preview['dates'],images/(re.sub(r'[<>:"/\\|?*]','_',group['tt'])+'.png'))
    for row in ws:
        ws.row_dimensions[row[0].row].height=27
        for c in row:
            c.alignment=Alignment(vertical='center',horizontal='left' if c.column==1 else 'center')
            if not c.font.name:c.font=Font(name='Arial',size=11)
        c=row[3]
        if isinstance(c.value,(int,float)):
            c.number_format='0.00%';c.fill=PatternFill('solid',fgColor=detail_color(c.value).lstrip('#'))
        elif c.data_type=='f':c.number_format='0.00%'
    notes=list(dict.fromkeys([*[g['tt']+': '+g['dateNote'] for g in preview['groups']],*preview['warnings']]))
    ns=wb.create_sheet('Примечания');ns.sheet_view.showGridLines=False;ns.column_dimensions['A'].width=120
    ns.append(['Примечания к данным']);ns['A1'].font=Font(name='Arial',bold=True,size=15,color=BLUE)
    for note in notes:
        ns.append([note]);cell=ns.cell(ns.max_row,1);cell.font=Font(name='Arial',size=11);cell.alignment=Alignment(wrap_text=True,vertical='top');ns.row_dimensions[ns.max_row].height=max(36,18*(len(note)//110+1))
    wb.save(folder/'Отчёт.xlsx')
    (folder/'Примечания.txt').write_text('Примечания к данным\n\n'+'\n\n'.join(notes),encoding='utf-8')
    audit={k:v for k,v in preview.items() if k!='groups'}
    audit.update(created=dt.datetime.now().isoformat(timespec='seconds'),stores=len(preview['groups']),dailyRecords=preview['badDays'],appealRecords=0,weeklyPct=preview['late']/preview['orders'] if preview['orders'] else None,sources=data['sources'],details=preview['groups'],checks=['Дни выбраны по общему просроку ТТ от 6% включительно','Все показатели детализации и итоги взяты из выгрузки сборщиков','Сверка с ТТ выполнена; расхождения отражены предупреждениями без изменения данных','Дедупликация заказов по сборочным листам не выполнялась: в выгрузке нет номеров заказов и листов','Сборщики без ФИО сохранены в итогах, кроме скрытого ID 0','Один PNG на ТТ, период до семи календарных дней','Итоги относятся только к выбранным дням'])
    (folder/'Проверка.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
    with zipfile.ZipFile(folder/'Все результаты.zip','w',zipfile.ZIP_DEFLATED) as archive:
        for p in [folder/'Отчёт.xlsx',folder/'Примечания.txt',folder/'Проверка.json',*sorted(images.glob('*.png'))]:archive.write(p,p.relative_to(folder))
    return audit
