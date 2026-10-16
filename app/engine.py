from __future__ import annotations
import collections, datetime as dt, hashlib, io, json, math, re, textwrap, zipfile
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from PIL import Image, ImageDraw, ImageFont

DEFAULT_CODES = ['1996','5494','1600','7261','3454','5607','5359','4708','6929','3022','2592','7717','5895','2719']
BLUE='294B63'; GREEN='C6EFCE'; LIGHT_RED='FCE0E0'; RED='F08080'
def norm(s):
    return re.sub(r'[^а-яa-z0-9]+','',str(s or '').casefold().replace('ё','е'))
def code(s):
    m=re.match(r'\s*(\d{4})',str(s or ''))
    return m[1] if m else str(s or '').strip().casefold()
def clean_name(s):
    return re.sub(r'\s+\d+\s*$','',re.sub(r'\s*\([^)]*\)','',str(s or ''))).strip()
def unknown(s):
    return not s or norm(s) in {'неопределенныйпользователь','фионеуказано','неизвестно'}
def number(x):
    if isinstance(x,(int,float)) and not isinstance(x,bool) and math.isfinite(x): return x
    if isinstance(x,str):
        try:
            v=float(x.replace('\xa0','').replace(' ','').replace(',','.'))
            if math.isfinite(v): return v
        except ValueError: pass
    raise ValueError(f'Ожидалось число, получено «{x}».')
def date(x,epoch):
    if isinstance(x,dt.datetime): return x.date().isoformat()
    if isinstance(x,dt.date): return x.isoformat()
    if isinstance(x,(float,int)):
        return openpyxl.utils.datetime.from_excel(x,epoch).date().isoformat()
    for fmt in ('%d.%m.%Y','%d.%m.%y','%Y-%m-%d','%d.%m.%Y %H:%M:%S'):
        try:return dt.datetime.strptime(str(x).strip(),fmt).date().isoformat()
        except ValueError:pass
    raise ValueError(f'Не удалось прочитать дату «{x}».')
ALIASES={
 'tt':['Наименование ТТ','ТТ'], 'date':['Дата и время начала слота','Дни','Дата'],
 'orders':['Заказы в статусе Успешно завершен','Кол-во заказов УЗ','Кол-во заказов (УЗ)'],
 'late':['Общее кол-во заказов, просроченных по сборке','Общее количество просроченных заказов'],
 'pct':['Общая доля заказов, просроченных по сборке','Общая доля(по дням)'],
 'helper':['Помощник','Лидер сервиса офлайн'], 'name':['Сборщик'], 'count':['абс','Количество обращений']}
def columns(row):
    hs={norm(v):i for i,v in enumerate(row) if v is not None}
    return {k:next((hs[norm(a)] for a in aliases if norm(a) in hs),None) for k,aliases in ALIASES.items()}

def analyze(files):
    daily={}; appeals=[]; warnings=[]; errors=[]; sources=[]; hashes=set(); ignored=0
    for filename,blob in files:
        digest=hashlib.sha256(blob).hexdigest()
        if digest in hashes:
            warnings.append(f'Повторная копия файла «{filename}» пропущена.');continue
        hashes.add(digest)
        try: book=openpyxl.load_workbook(io.BytesIO(blob),data_only=True,read_only=True)
        except Exception as e: raise ValueError(f'Не удалось открыть «{filename}». Нужен исправный XLSX.') from e
        recognized=False
        for sheet in book:
            rows=sheet.iter_rows(values_only=True); cs=None; kind=None
            for rn,row in enumerate(rows,1):
                if cs is None:
                    if rn>40:break
                    candidate=columns(row)
                    if all(candidate[k] is not None for k in ['tt','date','orders','late']):cs=candidate;kind='daily'
                    elif all(candidate[k] is not None for k in ['tt','date','name','count']):cs=candidate;kind='appeals'
                    if cs is not None:recognized=True;sources.append({'file':filename,'sheet':sheet.title,'kind':kind,'header':rn})
                    continue
                if not any(v is not None for v in row):continue
                get=lambda k:row[cs[k]] if cs.get(k) is not None and cs[k]<len(row) else None
                # Export footer totals are not records: no store and/or no date.
                if get('tt') is None or get('date') is None:
                    ignored+=1;continue
                try:
                    tt=str(get('tt')).strip(); day=date(get('date'),book.epoch); key=(code(tt),day)
                    if not re.match(r'^\d{4}',tt):raise ValueError('У точки нет четырёхзначного кода.')
                    if kind=='daily':
                        orders=number(get('orders'));late=number(get('late'))
                        if orders<0 or late<0 or late>orders:raise ValueError('Количество просроченных должно быть от 0 до числа заказов.')
                        pct=get('pct')
                        pct=number(pct) if pct is not None else (late/orders if orders else None)
                        if pct is not None and not 0<=pct<=1:raise ValueError('Доля просрока должна быть от 0 до 100%.')
                        rec={'tt':tt,'code':key[0],'date':day,'orders':orders,'late':late,'pct':pct,'helper':str(get('helper') or '').strip(),'source':f'{filename} / {sheet.title}, строка {rn}'}
                        if key in daily:
                            prev=daily[key]
                            if all(prev[x]==rec[x] for x in ['orders','late','pct']):
                                warnings.append(f'Повтор: {tt}, {day}; данные учтены один раз.');continue
                            errors.append(f'Разные значения для {tt}, {day}: {prev["source"]} и {rec["source"]}.');continue
                        daily[key]=rec
                    else:
                        count=number(get('count'))
                        if count<0:raise ValueError('Количество обращений не может быть отрицательным.')
                        name=clean_name(get('name'))
                        appeals.append({'tt':tt,'code':key[0],'date':day,'name':name,'count':count,'unknown':unknown(name)})
                except (ValueError,TypeError) as e:errors.append(f'{filename}, лист «{sheet.title}», строка {rn}: {e}')
        book.close()
        if not recognized:errors.append(f'В «{filename}» не найдены исходные столбцы ТТ, даты и показателей. Загружайте исходную выгрузку, а не готовый отчёт.')
    tts={}
    for r in [*appeals,*daily.values()]:tts[r['code']]={'code':r['code'],'name':r['tt']}
    if ignored:warnings.append(f'Служебных строк без ТТ или даты пропущено: {ignored}.')
    return {'daily':list(daily.values()),'appeals':appeals,'warnings':list(dict.fromkeys(warnings)), 'errors':errors,'sources':sources,'tts':sorted(tts.values(),key=lambda x:x['code']), 'unknownCount':sum(a['unknown'] for a in appeals)}

def select(data,opts):
    selected=set(opts.get('codes',DEFAULT_CODES)); exclude=opts.get('excludeUnknown',True)
    days=[r for r in data['daily'] if r['code'] in selected]
    apps=[r for r in data['appeals'] if r['code'] in selected and not(exclude and r['unknown'])] if opts.get('mode')!='late' else []
    dates={r['date'] for r in [*days,*apps]}
    if not dates:raise ValueError('Нет данных по выбранным торговым точкам.')
    lookup=collections.defaultdict(list)
    for a in apps:lookup[a['code'],a['date']].append(a)
    groups=[];consumed=0
    for c in sorted(selected):
        rs=sorted([r for r in days if r['code']==c],key=lambda r:r['date'])
        aa=[a for a in apps if a['code']==c]
        if not rs and not aa:continue
        tt=(rs or aa)[0]['tt'];matched=set(r['date'] for r in rs)
        for a in aa:
            if rs and a['date'] not in matched:raise ValueError(f'Обращение {tt} за {a["date"]} не имеет строки просрока. Добавьте выгрузку за этот день.')
        groups.append({'code':c,'tt':tt,'days':rs,'appeals':aa,'orders':sum(r['orders'] for r in rs),'late':sum(r['late'] for r in rs)})
        consumed+=len(aa)
    assert consumed==len(apps)
    return groups,sorted(dates),apps

def font(size=20,bold=False):
    path=Path('C:/Windows/Fonts')/('arialbd.ttf' if bold else 'arial.ttf')
    if not path.exists():raise ValueError('Для изображений нужен установленный шрифт Arial.')
    return ImageFont.truetype(str(path),size)
def color(v,good,bad):return '#'+(GREEN if v is not None and v<=good else RED if v is not None and v>bad else LIGHT_RED) if v is not None else '#EDF2F6'
def pct(v):return '—' if v is None else f'{v*100:.2f}%'.replace('.',',')
def n(v):return f'{v:g}' if v is not None else '—'

def picture(group,dates,path,good,bad,only_appeals=False):
    width=1500; x=[0,650,850,1050,1280,1500]; rows=[]
    amap=collections.defaultdict(list)
    for a in group['appeals']:amap[a['date']].append(a)
    if only_appeals:
        for day in sorted(amap):
            rows.append(([dt.date.fromisoformat(day).strftime('%d.%m.%Y'),'','','',''],None,'day'))
            for a in amap[day]:rows.append(([a['name'] or 'ФИО не указано','','','',n(a['count'])],None,'appeal'))
    else:
        for r in group['days']:
            rows.append(([dt.date.fromisoformat(r['date']).strftime('%d.%m.%Y'),n(r['orders']),n(r['late']),pct(r['pct']), ''],r['pct'],'day'))
            for a in amap[r['date']]:rows.append(([a['name'] or 'ФИО не указано','','','',n(a['count'])],None,'appeal'))
    ratio=group['late']/group['orders'] if group['orders'] else None
    rows.append((['Итого за неделю',n(group['orders']) if not only_appeals else '',n(group['late']) if not only_appeals else '',pct(ratio) if not only_appeals else '',n(sum(a['count'] for a in group['appeals']))],ratio,'total'))
    height=180+len(rows)*52;im=Image.new('RGB',(width,height),'white');draw=ImageDraw.Draw(im)
    draw.text((20,18),group['tt'],font=font(30,True),fill='#243E52')
    period=f'{dt.date.fromisoformat(dates[0]):%d.%m.%Y} — {dt.date.fromisoformat(dates[-1]):%d.%m.%Y}'
    draw.text((20,68),period,font=font(20),fill='#607484')
    draw.rectangle((0,115,width,180),fill='#294B63')
    labels=['Дата / ФИО','Заказы УЗ, шт.','Просрочено, шт.','Общий просрок','Обращения, шт.'] if not only_appeals else ['Дата / ФИО','','','','Обращения, шт.']
    for i,label in enumerate(labels):draw.text(((x[i]+x[i+1])/2,147),label,font=font(20,True),fill='white',anchor='mm')
    for j,(vals,v,kind) in enumerate(rows):
        y=180+j*52;fill='#DCE6ED' if kind=='total' else '#EFF4F7' if kind=='day' else 'white';draw.rectangle((0,y,width,y+52),fill=fill)
        if kind!='appeal' and not only_appeals:draw.rectangle((x[3],y,x[4],y+52),fill=color(v,good,bad))
        for i,value in enumerate(vals):
            f=font(21,kind=='total' or i==3)
            if i==0:
                while draw.textlength(value,font=f)>615 and f.size>14:f=font(f.size-1,kind=='total')
                draw.text((20,y+26),value,font=f,fill='#243D50',anchor='lm')
            else:draw.text(((x[i]+x[i+1])/2,y+26),value,font=f,fill='#243D50',anchor='mm')
        draw.line((0,y+51,width,y+51),fill='#DCE4EA')
    im.save(path)

def generate(data,opts,folder):
    if data['errors']:raise ValueError('Сначала устраните ошибки в загруженных данных.')
    groups,dates,apps=select(data,opts); good=float(opts.get('good',6))/100;bad=float(opts.get('bad',10))/100
    if not 0<=good<bad<=1:raise ValueError('Пороги должны быть от 0 до 100%, зелёный порог — меньше красного.')
    mode=opts.get('mode','combined');only_apps=mode=='appeals'
    if not only_apps and any(not g['days'] for g in groups):raise ValueError('Для отчёта просрока загрузите выгрузку по магазинам и дням.')
    if mode=='late':
        for g in groups:g['appeals']=[]
        apps=[]
    folder=Path(folder);folder.mkdir(parents=True,exist_ok=True);images=folder/'Изображения';images.mkdir(exist_ok=True)
    wb=openpyxl.Workbook();ws=wb.active;ws.title='По торговым точкам';ws.sheet_view.showGridLines=False
    ws.column_dimensions['A'].width=58
    for c in 'BCDE':ws.column_dimensions[c].width=21
    date_refs={};verify=[];row=1
    for g in groups:
        ws.cell(row,1,g['tt']);ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=5);ws.cell(row,1).font=Font(name='Arial',size=15,bold=True,color=BLUE);ws.row_dimensions[row].height=32;row+=1
        ws.cell(row,1,f'{dt.date.fromisoformat(dates[0]):%d.%m.%Y} — {dt.date.fromisoformat(dates[-1]):%d.%m.%Y}');ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=5);row+=1
        heads=['Дата / ФИО','Заказы УЗ, шт.','Просрочено, шт.','Общий просрок','Обращения, шт.']
        if only_apps:heads=['Дата / ФИО','','','','Обращения, шт.']
        for c,h in enumerate(heads,1):ws.cell(row,c,h);ws.cell(row,c).fill=PatternFill('solid',fgColor=BLUE);ws.cell(row,c).font=Font(name='Arial',bold=True,color='FFFFFF',size=10);ws.cell(row,c).alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        ws.row_dimensions[row].height=35;row+=1;first=row;dailyrows=[]
        amap=collections.defaultdict(list)
        for a in g['appeals']:amap[a['date']].append(a)
        records=g['days'] if not only_apps else [{'date':day} for day in sorted(amap)]
        for rec in records:
            vals=[dt.date.fromisoformat(rec['date']),rec.get('orders'),rec.get('late'),rec.get('pct'),None]
            for c,val in enumerate(vals,1):ws.cell(row,c,val);ws.cell(row,c).fill=PatternFill('solid',fgColor='EFF4F7')
            ws.cell(row,1).number_format='dd.mm.yyyy';ws.cell(row,4).number_format='0.00%';dailyrows.append(row);date_refs[g['code'],rec['date']]=row
            if not only_apps:verify.append((row,rec))
            row+=1
            for a in amap[rec['date']]:ws.cell(row,1,a['name'] or 'ФИО не указано');ws.cell(row,5,a['count']);row+=1
        total=row;ws.cell(row,1,'Итого за неделю')
        if not only_apps:
            for col in 'BC':ws[f'{col}{row}']='=SUM('+','.join(col+str(r) for r in dailyrows)+')'
            ws[f'D{row}']=f'=IF(B{row}=0,"",C{row}/B{row})';ws[f'D{row}'].number_format='0.00%'
        ws[f'E{row}']=f'=SUM(E{first}:E{row-1})'
        for c in ws[row]:c.fill=PatternFill('solid',fgColor='DCE6ED');c.font=Font(name='Arial',bold=True,size=11)
        if not only_apps:
            for rn in dailyrows+[row]:
                addr=f'D{rn}';ws.conditional_formatting.add(addr,CellIsRule(operator='lessThanOrEqual',formula=[good],fill=PatternFill('solid',fgColor=GREEN)))
                ws.conditional_formatting.add(addr,FormulaRule(formula=[f'AND(ISNUMBER({addr}),{addr}>{good},{addr}<={bad})'],fill=PatternFill('solid',fgColor=LIGHT_RED)))
                ws.conditional_formatting.add(addr,CellIsRule(operator='greaterThan',formula=[bad],fill=PatternFill('solid',fgColor=RED)))
        row+=3
        picture(g,dates,images/(re.sub(r'[<>:"/\\|?*]','_',g['tt'])+'.png'),good,bad,only_apps)
    for rows in ws:
        for c in rows:
            if not c.font.name:c.font=Font(name='Arial',size=11)
            if c.column>1:c.alignment=Alignment(horizontal='center',vertical='center')
        if ws.row_dimensions[rows[0].row].height is None:ws.row_dimensions[rows[0].row].height=23
    if opts.get('journal',True) and apps:
        j=wb.create_sheet('Журнал обращений');j.sheet_view.showGridLines=False;j.append(['Торговая точка',*map(dt.date.fromisoformat,dates)]);j.freeze_panes='B2';j.column_dimensions['A'].width=46
        for i in range(2,len(dates)+2):j.column_dimensions[openpyxl.utils.get_column_letter(i)].width=15;j.cell(1,i).number_format='dd.mm.yyyy'
        for g in groups:
            sums=collections.defaultdict(float)
            for a in g['appeals']:sums[a['date']]+=a['count']
            j.append([g['tt'],*[sums[d] if d in sums else None for d in dates]])
        for c in j[1]:c.fill=PatternFill('solid',fgColor=BLUE);c.font=Font(name='Arial',color='FFFFFF',bold=True)
    if opts.get('outline',False) and data['daily']:
        o=wb.create_sheet('Структура');o.sheet_properties.outlinePr.summaryBelow=False;o.append(['Помощник / ТТ / Дата','Заказы УЗ','Просрочено','Доля просрока']);helpers=collections.defaultdict(list)
        for g in groups:helpers[(g['days'][0].get('helper') or 'Помощник не указан') if g['days'] else 'Помощник не указан'].append(g)
        for helper,gg in sorted(helpers.items()):
            o.append([helper]);hr=o.max_row
            for g in gg:
                o.append([g['tt'],g['orders'],g['late'],g['late']/g['orders'] if g['orders'] else None]);tr=o.max_row;o.row_dimensions[tr].outlineLevel=1
                for r in g['days']:o.append([dt.date.fromisoformat(r['date']),r['orders'],r['late'],r['pct']]);o.row_dimensions[o.max_row].outlineLevel=2;o.cell(o.max_row,1).number_format='dd.mm.yyyy'
            for rn in range(hr+1,o.max_row+1):o.row_dimensions[rn].hidden=False
        o.column_dimensions['A'].width=50
        for col in 'BCD':o.column_dimensions[col].width=18
        for rn in range(2,o.max_row+1):o.cell(rn,4).number_format='0.00%'
    path=folder/'Отчёт.xlsx';wb.save(path)
    check=openpyxl.load_workbook(path,data_only=False)
    for rn,rec in verify:
        assert check.worksheets[0].cell(rn,2).value==rec['orders'] and check.worksheets[0].cell(rn,3).value==rec['late']
    assert sum(len(g['appeals']) for g in groups)==len(apps)
    assert len(list(images.glob('*.png')))==len(groups)
    totalorders=sum(g['orders'] for g in groups);totallate=sum(g['late'] for g in groups)
    audit={'created':dt.datetime.now().isoformat(timespec='seconds'),'stores':len(groups),'dailyRecords':sum(len(g['days']) for g in groups),'appealRecords':len(apps),'appeals':sum(a['count'] for a in apps),'orders':totalorders,'late':totallate,'weeklyPct':totallate/totalorders if totalorders else None,'excludedUnknown':data['unknownCount'] if opts.get('excludeUnknown',True) else 0,'thresholds':{'good':good,'bad':bad},'sources':data['sources'],'warnings':data['warnings'],'checks':['Дневные количества совпадают с исходниками','Все выбранные обращения распределены по ТТ и датам','Недельные итоги считаются по суммам, а не средним процентам','Один PNG для каждой ТТ','XLSX повторно открывается без ошибок']}
    (folder/'Проверка.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
    archive=folder/'Все результаты.zip'
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
        for file in [path,folder/'Проверка.json',*images.glob('*.png')]:z.write(file,file.relative_to(folder))
    return audit
